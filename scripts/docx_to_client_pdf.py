from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    BaseDocTemplate, PageTemplate, Frame, Paragraph as PdfParagraph, Spacer,
    Table, TableStyle, PageBreak, KeepTogether
)

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "client-manual" / "Swift-Financial-Client-User-Guide.docx"
OUT = ROOT / "docs" / "client-manual" / "Swift-Financial-Client-User-Guide.pdf"

INDIGO = colors.HexColor("#312E81")
INDIGO_2 = colors.HexColor("#3730A3")
INK = colors.HexColor("#1F2937")
MUTED = colors.HexColor("#6B7280")
LIGHT = colors.HexColor("#EEF2FF")

styles = getSampleStyleSheet()
body = ParagraphStyle("Body", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.4, leading=12.1, textColor=INK, spaceAfter=5)
h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=16, leading=19, textColor=INDIGO, spaceBefore=11, spaceAfter=6, keepWithNext=True)
h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=14, textColor=INDIGO_2, spaceBefore=8, spaceAfter=4, keepWithNext=True)
h3 = ParagraphStyle("H3", parent=styles["Heading3"], fontName="Helvetica-Bold", fontSize=10, leading=12, textColor=INK, spaceBefore=6, spaceAfter=3, keepWithNext=True)
bullet = ParagraphStyle("Bullet", parent=body, leftIndent=18, firstLineIndent=-9, bulletIndent=5, spaceAfter=3)
small = ParagraphStyle("Small", parent=body, fontSize=8.2, leading=10.2)
cover_brand = ParagraphStyle("CoverBrand", parent=body, fontName="Helvetica-Bold", fontSize=12, textColor=colors.HexColor("#C7D2FE"), spaceAfter=28)
cover_title = ParagraphStyle("CoverTitle", parent=body, fontName="Helvetica-Bold", fontSize=29, leading=34, textColor=colors.white, spaceAfter=12)
cover_sub = ParagraphStyle("CoverSub", parent=body, fontName="Helvetica", fontSize=13, leading=17, textColor=colors.HexColor("#E0E7FF"), spaceAfter=34)
cover_meta = ParagraphStyle("CoverMeta", parent=body, fontName="Helvetica", fontSize=9.5, leading=13, textColor=colors.HexColor("#C7D2FE"))

def header_footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica-Bold", 7.5)
    canvas.setFillColor(MUTED)
    canvas.drawString(0.8*inch, 10.58*inch, "SWIFT FINANCIAL  |  CLIENT USER GUIDE")
    canvas.setStrokeColor(colors.HexColor("#D1D5DB"))
    canvas.line(0.8*inch, 10.48*inch, 7.7*inch, 10.48*inch)
    canvas.setFont("Helvetica", 7.5)
    canvas.drawString(0.8*inch, 0.38*inch, "Client copy  |  Version 1.0  |  September 2026")
    canvas.drawRightString(7.7*inch, 0.38*inch, f"Page {doc.page}")
    canvas.restoreState()

def iter_blocks(parent):
    for child in parent.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield DocxTable(child, parent)

def clean(text):
    return escape(" ".join((text or "").replace("\u2014", " - ").replace("\u2013", "-").split()))

def make_table(block):
    data = []
    for row in block.rows:
        vals = [PdfParagraph(clean(cell.text), small) for cell in row.cells]
        data.append(vals)
    n = len(data[0]) if data else 1
    widths = {
        1: [6.72*inch], 2: [1.8*inch, 4.92*inch], 3: [1.3*inch, 3.45*inch, 1.97*inch],
        4: [1.1*inch, 1.8*inch, 2.2*inch, 1.62*inch]
    }.get(n, [6.72*inch/n]*n)
    t = Table(data, colWidths=widths, repeatRows=1 if len(data) > 1 else 0, hAlign="CENTER")
    if len(data) == 1:
        bg = LIGHT
        text_color = INK
    else:
        bg = INDIGO
        text_color = colors.white
    commands = [
        ("BACKGROUND", (0,0), (-1,0), bg), ("TEXTCOLOR", (0,0), (-1,0), text_color),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"), ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("GRID", (0,0), (-1,-1), 0.45, colors.HexColor("#D1D5DB")),
        ("LEFTPADDING", (0,0), (-1,-1), 7), ("RIGHTPADDING", (0,0), (-1,-1), 7),
        ("TOPPADDING", (0,0), (-1,-1), 6), ("BOTTOMPADDING", (0,0), (-1,-1), 6),
    ]
    if len(data) > 2:
        for r in range(2, len(data), 2): commands.append(("BACKGROUND", (0,r), (-1,r), colors.HexColor("#F9FAFB")))
    t.setStyle(TableStyle(commands))
    return t

source = Document(SRC)
story = []
first_table = True
previous_h1 = False
num = 0
for block in iter_blocks(source):
    if isinstance(block, DocxTable):
        if first_table:
            first_table = False
            cell_text = [p.text for p in block.cell(0,0).paragraphs]
            cover_flow = [
                PdfParagraph(clean(cell_text[0] if cell_text else "SWIFT FINANCIAL"), cover_brand),
                PdfParagraph(clean(cell_text[1] if len(cell_text)>1 else "Client User Guide"), cover_title),
                PdfParagraph(clean(cell_text[2] if len(cell_text)>2 else "A practical guide"), cover_sub),
                PdfParagraph(clean(cell_text[3] if len(cell_text)>3 else "Version 1.0"), cover_meta),
            ]
            ctab = Table([[cover_flow]], colWidths=[6.72*inch], rowHeights=[5.1*inch])
            ctab.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),INDIGO),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("LEFTPADDING",(0,0),(-1,-1),28),("RIGHTPADDING",(0,0),(-1,-1),28)]))
            story += [Spacer(1, 0.65*inch), ctab, Spacer(1, 0.32*inch)]
        else:
            story += [make_table(block), Spacer(1, 4)]
        continue
    text = block.text.strip()
    if not text:
        continue
    style_name = block.style.name if block.style else ""
    if style_name == "Heading 1":
        num = 0
        if text != "Document control": story.append(PageBreak())
        story.append(PdfParagraph(clean(text), h1))
        previous_h1 = True
    elif style_name == "Heading 2":
        num = 0
        story.append(PdfParagraph(clean(text), h2))
    elif style_name == "Heading 3":
        num = 0
        story.append(PdfParagraph(clean(text), h3))
    elif style_name == "List Bullet": story.append(PdfParagraph(clean(text), bullet, bulletText="•"))
    elif style_name == "List Number":
        num += 1
        story.append(PdfParagraph(clean(text), bullet, bulletText=f"{num}."))
    else:
        if not text[0].isdigit(): num = 0
        story.append(PdfParagraph(clean(text), body))

doc = BaseDocTemplate(str(OUT), pagesize=letter, leftMargin=0.8*inch, rightMargin=0.8*inch, topMargin=0.62*inch, bottomMargin=0.62*inch, title="Swift Financial Client User Guide", author="Swift Financial")
frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="normal")
doc.addPageTemplates([PageTemplate(id="main", frames=[frame], onPage=header_footer)])
doc.build(story)
print(OUT)
